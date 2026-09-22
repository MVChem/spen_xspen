from pathlib import Path
import math
import re
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

HERE=Path(__file__).resolve().parent
OUT=HERE.parents[1]/'ppt'
OUT.mkdir(parents=True, exist_ok=True)
ASSETS=HERE/'assets'
R=Presentation(); R.slide_width=Inches(15.86); R.slide_height=Inches(9.92)
S=R.slides.add_slide(R.slide_layouts[6])
NAVY='071D43'; BLUE='0069CF'; TEAL='00A8AA'; PURPLE='7928BA'; ORANGE='F87C00'
def u(p): return Inches(p/100)
def rgb(c): return RGBColor.from_string(c)
def box(name,x,y,w,h,fill=None,line=None,width=1.5,round=False,dash=False):
    sh=S.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,u(x),u(y),u(w),u(h)); sh.name=name
    if round: sh.adjustments[0]=0.015
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb=rgb(line); sh.line.width=Pt(width*.72)
    else: sh.line.fill.background()
    if dash:
        el=OxmlElement('a:prstDash'); el.set('val','dash'); sh.line._get_or_add_ln().append(el)
    return sh

def text(name,txt,x,y,w,h,size=16,bold=False,color=NAVY,align='center',font='Arial',italic=False):
    sh=S.shapes.add_textbox(u(x),u(y),u(w),u(h)); sh.name=name
    tf=sh.text_frame; tf.clear(); tf.word_wrap=False; tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment={'center':PP_ALIGN.CENTER,'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT}[align]
        p.space_before=Pt(0); p.space_after=Pt(0)
        run=p.add_run(); run.text=line; run.font.name=font; run.font.size=Pt(size*.72); run.font.bold=bold; run.font.italic=italic; run.font.color.rgb=rgb(color)
    return sh

def eq(name,txt,x,y,w,h,size=20):
    replacements={'xσ':'x_{σ}','Dθ':'D_{θ}','ρσ':'ρ_{σ}','x₀':'x_{0}','x̂₀':'x̂_{0}','xᴰᶜ':'x_{DC}','minₓ':'min_{x}','xₜ₋₁':'x_{t−1}','xₜ':'x_{T}','z꜀':'z_{c}','Ŝ꜀':'Ŝ_{c}','F꜀':'F_{c}','Fφ,꜀':'F_{φ,c}','Aᴴ':'A^{H}','Dᴿᴼ':'D_{RO}','Aᴾᴱ':'A_{PE}','kᴿᴼ':'k_{RO}','²':'^{2}','ᴺᶜ × ᴺᴾᴱ × ᴺᴿᴼ':'^{N꜀ × N_{PE} × N_{RO}}'}
    for a,b in replacements.items(): txt=txt.replace(a,b)
    txt=txt.replace('^{N꜀ × N_{PE} × N_{RO}}','^{Nc × NPE × NRO}')
    sh=text(name,'',x,y,w,h,size,font='Liberation Serif',italic=True)
    p=sh.text_frame.paragraphs[0]
    for token in re.split(r'([_^]\{[^}]*\})',txt):
        if not token: continue
        special=token.startswith(('_{','^{'))
        r=p.add_run(); r.text=token[2:-1] if special else token
        r.font.name='Liberation Serif'; r.font.size=Pt(size*.72*(.7 if special else 1)); r.font.italic=True; r.font.color.rgb=rgb(NAVY)
        if special: r._r.get_or_add_rPr().set('baseline','-22000' if token[0]=='_' else '35000')
    return sh
def path(name,pts,color=BLUE,width=2,arrow=False,dash=False):
    b=S.shapes.build_freeform(pts[0][0],pts[0][1],scale=u(1)); b.add_line_segments(pts[1:],close=False); sh=b.convert_to_shape(); sh.name=name
    sh._element.spPr.append(OxmlElement('a:effectLst')); sh.fill.background(); sh.line.color.rgb=rgb(color); sh.line.width=Pt(width*.72)
    ln=sh.line._get_or_add_ln()
    if dash:
        e=OxmlElement('a:prstDash'); e.set('val','dash'); ln.append(e)
    if arrow:
        e=OxmlElement('a:tailEnd'); e.set('type','triangle'); e.set('w','med'); e.set('len','med'); ln.append(e)
    return sh

def poly(name,pts,fill,line=BLUE):
    b=S.shapes.build_freeform(*pts[0],scale=u(1)); b.add_line_segments(pts[1:],close=True); sh=b.convert_to_shape(); sh.name=name
    sh._element.spPr.append(OxmlElement('a:effectLst')); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill); sh.line.color.rgb=rgb(line); sh.line.width=Pt(1.2); return sh

im=Image.open(ASSETS/'reference.png')
def crop(name,rect):
    fp=ASSETS/(name+'.png'); im.crop(rect).save(fp)
    x,y,x2,y2=rect; sh=S.shapes.add_picture(str(fp),u(x),u(y),u(x2-x),u(y2-y)); sh.name=name; return sh

# Panel structure and headings.
box('Training panel',20,60,1546,200,'EFF9FE','298FDC',3,True)
box('Sampling panel',20,276,1546,663,'FFFFFF','006AAA',3,True)
text('Figure title','SPEN Diffusion Reconstruction',250,0,1080,52,44,True)
text('Panel a','(a)  Training the magnitude prior',32,65,760,31,26,True,align='left')
text('Training mode','Unconditional EDM training',1240,67,310,24,16,False,align='right')
text('Panel b','(b)  Physics-constrained sampling',32,280,820,32,26,True,align='left')
# Training row
crop('Clean training magnitude',(106,123,230,219)); text('Clean image label','Clean magnitude image',72,101,190,21,14,True); eq('x0','x₀',127,217,80,29,22)
box('Gaussian noising',290,138,199,68,'E1F3FF','286BC2',2.4,True)
text('Noise instruction','Add Gaussian noise',300,148,179,24,17,True); eq('Noising equation','xσ = x₀ + σ ε',303,174,175,25,20)
eq('Noise distribution','ε ∼ N(0, I)',307,211,164,25,18)
crop('Noisy training image',(537,123,654,219)); text('Noisy image label','Noisy image',527,101,137,21,15,True); eq('Noisy symbol','xσ',553,219,85,25,21)
for i,pts in enumerate([[(746,119),(776,130),(776,216),(746,227)],[(783,135),(797,141),(797,205),(783,211)],[(921,141),(936,135),(936,211),(921,205)],[(943,130),(973,119),(973,227),(943,216)]]): poly('U-Net block '+str(i+1),pts,'8FCFFF')
box('U-Net central block',803,145,113,57,'B3DFFF',BLUE,2)
eq('Trainable denoiser','Dθ(xσ, σ)',805,156,108,32,22); text('U-Net title','U-Net denoiser',765,100,190,22,15,True)
crop('Training denoised output',(1037,123,1159,220)); text('Denoised output label','Denoised estimate',1013,102,172,22,14,True); eq('Denoised symbol','Dθ(xσ, σ)',1030,220,141,27,20)
box('Training loss box',1210,140,304,77,'EAF7FF','3577BE',2,True)
text('Loss label','Training loss (clean target)',1220,147,285,23,15,True)
eq('Training loss','ℒ = E[ w(σ) ‖Dθ(xσ, σ) − x₀‖² ]',1222,174,280,31,20)
for i,(a,b) in enumerate([(234,282),(492,534),(659,738),(977,1033),(1162,1209)]): path('Training arrow '+str(i),[(a,174),(b,174)],arrow=True)
# Calibration strip.
text('Measured title','Measured complex multi-coil SPEN data',65,316,296,22,15.5,True); eq('Observation symbol','y',357,314,30,24,23)
text('Measurement domain','After scanner phase correction and RO FFT',64,337,287,21,13)
crop('Measured complex coil data',(63,361,220,459)); eq('Measured data dimensions','y ∈ ℂᴺᶜ × ᴺᴾᴱ × ᴺᴿᴼ',64,461,167,25,18)
text('Acquisition grid label','Acquisition grid\n(e.g., 96 × 96)',238,359,114,35,13)
crop('Acquisition grid',(255,395,337,469))
box('Weighted adjoint box',531,376,125,61,'F7F0FE',PURPLE,1.6,True)
text('Weighted adjoint title','Weighted InvA',538,387,111,25,15,True); text('Weighted adjoint subtitle','Weighted adjoint',536,411,116,23,13,color=PURPLE)
eq('Weighted adjoint equation','z꜀ = Aᴴ W y',537,440,113,29,21)
crop('Complex coil images',(681,354,782,448)); text('Coil image title','Complex coil images',659,321,182,23,15,True); eq('Coil symbol','z꜀',706,446,50,27,20)
box('Coil normalization box',811,376,121,61,'F8F2FE',PURPLE,1.6,True)
eq('Coil normalization equation','Ŝ꜀ = z꜀ / RSS(z)',816,389,110,34,18)
text('Coil interpolation note','Interpolate and normalize\nfor finer grids',800,441,145,32,11.8)
box('Fixed maps boundary',952,347,256,126,None,PURPLE,1.5,True,True)
text('Fixed maps title','Fixed complex coil fractions: coil + object phase',890,306,339,22,13,True,PURPLE)
text('Fixed maps source','(derived from measurements)',944,325,232,19,13,color=PURPLE)
crop('Coil magnitude map',(963,376,1065,466)); crop('Coil phase map',(1086,376,1171,466)); crop('Coil phase colorbar',(1178,375,1203,468))
text('Map magnitude label','Magnitude |Ŝ꜀|',963,350,101,24,14,True); text('Map phase label','Phase ∠Ŝ꜀',1083,350,92,24,14,True)
for i,(a,b) in enumerate([(344,531),(658,680),(783,809),(934,951)]): path('Calibration arrow '+str(i),[(a,410),(b,410)],PURPLE,1.5,True)
# Freeze and measurement routing.
path('Frozen weights transfer',[(846,207),(846,289),(503,289),(503,497),(449,497),(449,545)],'063D92',3,True)
text('Freeze weights annotation','Freeze trained weights',859,237,176,23,14,True,'00589E')
path('Observation data consistency link',[(297,473),(297,483),(1131,483),(1131,473)],PURPLE,1.5)
path('Observation into consistency',[(503,497),(937,497),(937,547)],PURPLE,1.7,True); eq('y on observation branch','y',942,491,31,30,22)
# Main sampling row.
crop('Gaussian sampling initialization',(66,553,187,638)); text('Sampling initialization label','Gaussian initialization',41,509,170,22,14,True); eq('Sampling initialization distribution','xₜ ∼ N(0, I)',66,530,123,23,18)
box('Frozen denoiser box',344,545,217,77,'C5E5FF',BLUE,2.6,True)
text('Frozen denoiser label','Frozen denoiser',355,561,195,26,19,True); eq('Frozen denoiser formula','Dθ(x, σ)',387,585,132,29,24)
crop('Clean sampling estimate',(616,549,726,631)); text('Clean estimate label','Clean estimate',614,515,113,19,12,True); eq('Clean estimate symbol','x̂₀',638,530,64,19,20)
box('Data consistency box',830,548,207,59,'FFDCBA',ORANGE,3,True); text('Data consistency label','Data consistency',842,565,184,28,19,True)
eq('Data consistency objective','xᴰᶜ = arg minₓ ‖F(x) − y‖² + ρσ ‖x − x̂₀‖²',765,609,344,35,18)
box('Diffusion update box',1143,544,178,72,'C6E7FF',BLUE,2,True); text('Diffusion update label','Diffusion update',1152,559,160,25,17,True); eq('Next sample','xₜ₋₁',1190,582,84,26,20)
crop('Final reconstruction',(1401,548,1525,640)); text('Reconstruction label','Reconstruction',1405,506,116,23,14,True); eq('Reconstruction scaling','m = (x + 1)/2',1403,526,121,23,16)
for i,(a,b) in enumerate([(192,340),(565,614),(732,827),(1042,1140),(1326,1397)]): path('Sampling arrow '+str(i),[(a,584),(b,584)],BLUE,2,True)
path('Diffusion sampling loop',[(1168,616),(1168,661),(463,661),(463,626)],'0061BD',2.6,True)
box('Loop caption backing',686,647,160,26,'FFFFFF'); text('Loop caption','Repeat as σ decreases',690,649,150,20,12.5,True)
# Physics inset.
box('Physics inset',33,679,1139,242,'E9FAF8',TEAL,2.5,True)
text('Physics inset heading','What the SPEN degradation operator does',69,684,800,37,26,True,align='left')
box('Physics sequential operator boundary',191,728,796,151,'F1FCFC',TEAL,1.4,True)
text('Denoiser output label','Denoiser output',58,721,111,19,11.8,True); text('Denoiser output detail','(normalized real magnitude)',43,738,145,17,10.5)
crop('Physics input magnitude',(63,762,147,832)); eq('Physics input symbol','x',71,830,62,27,23)
box('Magnitude conversion',206,772,103,47,'EDFBFB',TEAL,1.5,True)
eq('Magnitude conversion equation','m = (x + 1)/2',213,782,90,25,17)
box('Coil multiplication',335,771,78,47,'F5FCFC',TEAL,1.5,True); eq('Coil multiplication formula','× Ŝ꜀',342,780,64,30,23)
text('Coil multiplication detail','Apply complex\ncoil sensitivities',325,821,106,30,11)
# Purple map route entering multiplication.
path('Fixed maps to physics',[(529,750),(529,723),(372,723),(372,770)],'9347CC',2,True)
box('SPEN encoding box',452,754,149,85,'F1FDFD',TEAL,1.4,True); text('SPEN encoding label','Aᴾᴱ: SPEN encoding',459,758,137,21,12,True)
crop('SPEN encoding matrix',(484,782,564,835)); text('SPEN encoding detail','Same scanner trajectory',455,841,145,22,11.5)
box('Readout bandlimit box',626,754,162,87,'F4FCFC',TEAL,1.4,True); text('Readout bandlimit label','Dᴿᴼ: RO band limitation',631,758,153,21,12,True)
box('Retained RO band',688,782,33,43,'C0E1E2')
path('Spectrum horizontal axis',[(632,824),(779,824)],NAVY,.8,True); path('Spectrum vertical axis',[(661,834),(661,784)],NAVY,.8,True)
vals=[0,1,0,2,0,3,0,4,1,8,4,2,10,4,12,8,16,12,24,20,30,15,27,12,18,5,14,4,8,2,5,1,3,0,2,1,0,0]
path('Illustrative RO spectrum',[(632+i*3.75,824-v) for i,v in enumerate(vals)],'222222',1.1)
eq('RO frequency label','kᴿᴼ',751,821,33,21,12); text('RO bandwidth detail','Retain acquired bandwidth',622,842,172,22,11.4)
box('PE sampling mask',811,754,170,85,'F1FCFC',TEAL,1.4,True); text('PE mask title','M: acquired PE rows',815,758,162,21,12,True)
for i in range(17): path('PE mask row '+str(i),[(819,782+i*3),(883,782+i*3)],BLUE if i<8 else 'ABC0CC',1.1,dash=i>=8)
path('Retained rows key',[(895,794),(906,794)],BLUE,3); path('Missing rows key',[(895,810),(906,810)],'A1AFBA',1.5,dash=True)
text('Retained rows legend','Retained rows',910,785,69,18,9.5,align='left'); text('Missing rows legend','Missing rows',910,802,69,18,9.5,align='left')
text('PE sampling detail','Full or undersampled PE',810,842,173,22,11.5)
box('Predicted data box',1008,727,152,152,'F5FDFD',TEAL,1.5,True)
text('Predicted data label','Predicted complex data',1016,734,138,20,12,True); crop('Predicted complex coil data',(1036,760,1128,826))
eq('Predicted data symbol','F꜀(x)',1030,826,109,24,19); eq('Predicted data dimensions','∈ ℂᴺᶜ × ᴺᴾᴱ × ᴺᴿᴼ',1018,848,132,24,14)
for i,(a,b) in enumerate([(152,193),(310,336),(414,452),(603,627),(790,812),(981,1008)]): path('Physics arrow '+str(i),[(a,795),(b,795)],TEAL,1.5,True)
path('Physics operator to consistency',[(620,728),(620,702),(906,702),(906,638)],TEAL,2,True)
eq('Forward operator equation','F꜀(x) = M Dᴿᴼ Aᴾᴱ [Ŝ꜀ (x + 1)/2]',416,879,376,35,21)
text('Grid size example','Example: 192 × 192 image grid → 96 × 96 acquisition',863,892,299,23,11.4)
# Experimental inset.
box('Phase refinement boundary',1245,679,311,236,None,PURPLE,1.4,True,True)
text('Phase refinement title','Experimental extension: phase refinement',1254,686,295,26,15,True,PURPLE)
text('Residual phase label','Residual phase φ',1267,716,111,23,12,True)
crop('Residual phase map',(1260,742,1338,829)); crop('Residual phase colorbar',(1346,741,1373,828))
text('Phase fitting note','•  Fit residual odd/even phase\n   from measured data',1381,737,171,35,11,align='left')
text('Phase row note','•  Even rows: exp(iφ); odd rows: 1',1381,774,172,23,10.6,align='left')
text('Phase updating note','•  Update phase during sampling;\n   prior stays frozen.',1381,799,171,33,10.8,align='left')
eq('Phase variant equation','Fφ,꜀(x) = M{exp(iEφ) ⊙ Dᴿᴼ Aᴾᴱ[Ŝ꜀m]}',1252,845,297,31,17.5)
text('Phase row placement','E: original even scanner rows, before PE masking',1254,879,292,25,11.5,color=PURPLE)
path('Optional phase connector',[(1243,794),(1177,794)],'A762DC',2,True,True); text('Optional connector label','Optional',1180,764,65,24,12,True,'9954D4')
text('Legend learned','Blue: learned prior',414,953,135,29,15,True,BLUE)
text('Legend physics','Teal: fixed physics',563,953,136,29,15,True,'00989E')
text('Legend phase','Purple: phase / coil estimates',708,953,198,29,15,True,PURPLE)
text('Legend consistency','Orange: measurement consistency',919,953,248,29,15,True,'F16800')
# Suppress theme-inherited shadows on all diagram elements.
for sh in S.shapes:
    for style in sh._element.findall('{http://schemas.openxmlformats.org/presentationml/2006/main}style'):
        sh._element.remove(style)
R.save(OUT/'fig1_v1.pptx')
print(OUT/'fig1_v1.pptx')
print('Objects:',len(S.shapes))
