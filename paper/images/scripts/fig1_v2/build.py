from pathlib import Path
import math
import re
from PIL import Image
from make_assets import make_assets, GENERATED
RENDER_META = make_assets()
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

HERE=Path(__file__).resolve().parent
OUT=HERE.parents[1]/'ppt'
OUT.mkdir(parents=True, exist_ok=True)
ASSETS=HERE/'assets'
R=Presentation(); R.slide_width=Inches(14.40); R.slide_height=Inches(9.01)
S=R.slides.add_slide(R.slide_layouts[6])
NAVY='071D43'; BLUE='146CA5'; TEAL='009C9D'; PURPLE='912CB1'; ORANGE='E7781B'
def u(p): return Inches(p/100)
def rgb(c): return RGBColor.from_string(c)
def box(name,x,y,w,h,fill=None,line=None,width=1.5,round=False,dash=False):
    sh=S.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,u(x),u(y),u(w),u(h)); sh.name=name
    if round: sh.adjustments[0]=0.015
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb=rgb(line); sh.line.width=Pt(width*.72)
    else: sh.line.fill.background()
    if dash:
        el=OxmlElement('a:prstDash'); el.set('val','dash'); sh.line._get_or_add_ln().append(el)
    return sh

def text(name,txt,x,y,w,h,size=16,bold=False,color=NAVY,align='center',font='Arial',italic=False):
    sh=S.shapes.add_textbox(u(x),u(y),u(w),u(h)); sh.name=name
    tf=sh.text_frame; tf.clear(); tf.word_wrap=False; tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment={'center':PP_ALIGN.CENTER,'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT}[align]
        p.space_before=Pt(0); p.space_after=Pt(0)
        run=p.add_run(); run.text=line; run.font.name=font; run.font.size=Pt(size*.72); run.font.bold=bold; run.font.italic=italic; run.font.color.rgb=rgb(color)
    return sh

def eq(name,txt,x,y,w,h,size=20):
    replacements={'xσ':'x_{σ}','Dθ':'D_{θ}','ρσ':'ρ_{σ}','x₀':'x_{0}','x̂₀':'x̂_{0}','xᴰᶜ':'x_{DC}','minₓ':'min_{x}','xₜ₋₁':'x_{t−1}','xₜ':'x_{T}','z꜀':'z_{c}','Ŝ꜀':'Ŝ_{c}','F꜀':'F_{c}','Fφ,꜀':'F_{φ,c}','Aᴴ':'A^{H}','Dᴿᴼ':'D_{RO}','Aᴾᴱ':'A_{PE}','kᴿᴼ':'k_{RO}','²':'^{2}','ᴺᶜ × ᴺᴾᴱ × ᴺᴿᴼ':'^{N꜀ × N_{PE} × N_{RO}}'}
    for a,b in replacements.items(): txt=txt.replace(a,b)
    txt=txt.replace('^{N꜀ × N_{PE} × N_{RO}}','^{Nc × NPE × NRO}')
    sh=text(name,'',x,y,w,h,size,font='Liberation Serif',italic=True)
    p=sh.text_frame.paragraphs[0]
    for token in re.split(r'([_^]\{[^}]*\})',txt):
        if not token: continue
        special=token.startswith(('_{','^{'))
        r=p.add_run(); r.text=token[2:-1] if special else token
        r.font.name='Liberation Serif'; r.font.size=Pt(size*.72*(.7 if special else 1)); r.font.italic=True; r.font.color.rgb=rgb(NAVY)
        if special: r._r.get_or_add_rPr().set('baseline','-22000' if token[0]=='_' else '35000')
    return sh
def path(name,pts,color=BLUE,width=2,arrow=False,dash=False):
    b=S.shapes.build_freeform(pts[0][0],pts[0][1],scale=u(1)); b.add_line_segments(pts[1:],close=False); sh=b.convert_to_shape(); sh.name=name
    sh._element.spPr.append(OxmlElement('a:effectLst')); sh.fill.background(); sh.line.color.rgb=rgb(color); sh.line.width=Pt(width*.72)
    ln=sh.line._get_or_add_ln()
    if dash:
        e=OxmlElement('a:prstDash'); e.set('val','dash'); ln.append(e)
    if arrow:
        e=OxmlElement('a:tailEnd'); e.set('type','triangle'); e.set('w','med'); e.set('len','med'); ln.append(e)
    return sh

def poly(name,pts,fill,line=BLUE):
    b=S.shapes.build_freeform(*pts[0],scale=u(1)); b.add_line_segments(pts[1:],close=True); sh=b.convert_to_shape(); sh.name=name
    sh._element.spPr.append(OxmlElement('a:effectLst')); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill); sh.line.color.rgb=rgb(line); sh.line.width=Pt(1.2); return sh

def picture(name,filename,rect):
    x,y,x2,y2=rect
    fp=GENERATED/(filename+'.png')
    iw,ih=Image.open(fp).size
    scale=min((x2-x)/iw,(y2-y)/ih)
    w,h=iw*scale,ih*scale
    sh=S.shapes.add_picture(str(fp),u(x+(x2-x-w)/2),u(y+(y2-y-h)/2),u(w),u(h));sh.name=name
    return sh

def image_stack(name,prefix,x,y,w,h,front=3):
    # Four separate, replaceable images, native outlines, one small semantic group.
    iw,ih=Image.open(GENERATED/(prefix+str(front)+'.png')).size
    scale=min(w/iw,h/ih); nw,nh=iw*scale,ih*scale
    x+=(w-nw)/2; y+=(h-nh)/2; w,h=nw,nh
    shapes=[]
    order=[c for c in range(4) if c!=front]+[front]
    for i,c in enumerate(order):
        offset=(3-i)*3
        shapes.append(picture(f'{name} coil {c}',prefix+str(c),(x+offset,y-offset,x+w+offset,y+h-offset)))
        shapes.append(box(f'{name} outline {c}',x+offset,y-offset,w,h,None,'37434A',.55))
    group=S.shapes.add_group_shape(shapes);group.name=name
    return group

def unet(name,x,y,w,h):
    shapes=[]
    for i,(dx,dy,bw,bh) in enumerate([(0,0,.11,1),(.14,.19,.1,.70),(.26,.37,.1,.4),(.63,.37,.1,.4),(.76,.19,.1,.70),(.89,0,.11,1)]):
        shapes.append(box(f'{name} block {i+1}',x+w*dx,y+h*dy,w*bw,h*bh,'54A4E1','3589BE',1))
    shapes.append(path(name+' outer skip',[(x+w*.17,y+h*.18),(x+w*.17,y+h*.09),(x+w*.85,y+h*.09)],BLUE,1.4,True))
    shapes.append(path(name+' inner skip',[(x+w*.31,y+h*.36),(x+w*.31,y+h*.28),(x+w*.71,y+h*.28)],BLUE,1.4,True))
    for j in range(3):
        sh=S.shapes.add_shape(MSO_SHAPE.OVAL,u(x+w*(.45+j*.046)),u(y+h*.57),u(3),u(3)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(BLUE); sh.line.fill.background(); sh.name=name+f' ellipsis {j+1}'; shapes.append(sh)
    g=S.shapes.add_group_shape(shapes); g.name=name

def fraction(name,numer,denom,x,y,w,size=18):
    eq(name+' numerator',numer,x,y,w,size+3,size)
    path(name+' fraction bar',[(x+2,y+size+3),(x+w-2,y+size+3)],NAVY,.8)
    eq(name+' denominator',denom,x,y+size+4,w,size+3,size)

# The two panels.
box('Training panel',10,56,1420,214,'F8FCFE','2385B5',2.7,True)
box('Sampling panel',10,300,1420,564,'FFFFFF','009593',2.7,True)
text('Figure title','SPEN Diffusion Reconstruction',200,2,1040,43,38,True,'13354C')
text('Panel a title','(a) Training the magnitude prior',20,62,750,33,25,True,'135F96',align='left')
text('Training mode','Unconditional EDM training',1110,68,300,25,17,color='155084',align='right')
text('Panel b title','(b) Physics-constrained sampling',20,306,780,31,24,True,'086E72',align='left')
# Training row.
text('Clean MRI title','Clean magnitude MRI',44,98,174,23,15,True); eq('Clean MRI symbol','x₀',108,116,38,23,20)
picture('Clean MRI — measured training slice','training_clean',(69,138,193,248))
text('Noise addition title','Add Gaussian noise',218,151,153,23,15)
eq('Noise addition formula','xσ = x₀ + σ ε',218,173,153,24,21)
path('Clean to noisy',[(200,197),(385,197)],BLUE,1.7,True)
text('Noisy MRI title','Noisy image',383,98,133,23,15); eq('Noisy MRI symbol','xσ',427,116,44,23,20)
picture('Noisy MRI — seeded Gaussian noising','training_noisy',(390,138,510,248)); path('Noisy to training U-Net',[(518,192),(592,192)],BLUE,1.7,True)
unet('Training U-Net',596,123,194,116); eq('Training U-Net formula','Dθ(xσ, σ)',631,93,121,29,23); text('Training U-Net label','U-Net',642,237,100,23,16,True,'135F96')
path('Training U-Net to output',[(795,192),(867,192)],BLUE,1.7,True)
text('Denoised estimate title','Denoised estimate',848,98,160,23,15); eq('Denoised estimate formula','Dθ(xσ, σ)',863,116,132,23,20)
picture('Denoised estimate — actual checkpoint output','training_denoised',(872,139,987,247))
path('Clean target loss link',[(996,195),(1125,195)],NAVY,1.3,True,True)
text('Loss link label','Loss\n(to clean target)',1003,153,119,35,13)
box('Training loss frame',1128,166,279,54,'F0F8FD',BLUE,1.4,True)
eq('Training loss formula','ℒ = E[ w(σ) ‖Dθ(xσ, σ) − x₀‖² ]',1135,174,265,38,21)
path('Freeze training weights',[(719,271),(719,298)],'008BA0',2.4,True)
text('Freeze weights label','Freeze trained weights',738,272,244,26,15,True,align='left')
# Measurement, acquisition axes, calibration.
text('Measured SPEN title','Measured SPEN data y\n(multi-coil, complex)',35,347,175,38,14.5)
image_stack('Measured SPEN data — four receiver coils','observation_coil',51,397,131,96)
path('Measured PE axis',[(41,499),(41,395)],'333333',.8,True)
sh=text('Measured PE axis label','PE (acquired rows)',-15,439,93,18,10.5); sh.rotation=270
path('Measured RO axis',[(50,505),(189,505)],'333333',.8,True)
text('Measured RO axis label','RO (after FFT)',59,505,129,18,11)
text('Scanner correction note','After scanner phase correction\nand RO FFT',24,521,192,35,12)
box('Calibration frame',218,343,444,202,'FAF3FD','B866CC',1.4,True,True)
text('Calibration title','Calibration from measurements\n(not ground-truth maps)',269,348,313,33,13,True,'7F2597')
box('Weighted adjoint frame',227,410,84,75,'F3E7F9',PURPLE,1.3,True)
text('Weighted adjoint label','Weighted\nInvA',232,422,74,35,13.5,True,'7F2597')
eq('Weighted adjoint description','(w^{†}, adjoint)',229,456,81,23,13)
path('Measured to calibration',[(199,447),(226,447)],PURPLE,1.7,True)
path('Weighted adjoint to coil images',[(312,449),(332,449)],PURPLE,1.4,True)
text('Complex coil images title','Complex\ncoil images',337,386,69,32,12,color='682285')
eq('Complex coil symbol','z꜀',348,415,47,24,21); image_stack('Native InvA coil images','coil_image',335,447,52,52)
text('Multiple coils ellipsis','…',397,455,15,24,16)
path('Coil images to normalization',[(400,451),(422,451)],PURPLE,1.2,True)
box('Normalization frame',424,423,72,56,'FCF6FF',PURPLE,1.2,True)
eq('Normalization left','Ŝ꜀ =',427,433,30,26,15); fraction('Normalization ratio','z꜀','RSS(z)',458,428,34,14)
text('Magnitude map title','Magnitude',508,383,62,19,11,color='682285'); eq('Magnitude map symbol','|Ŝ꜀|',515,399,50,21,18)
image_stack('Fixed 192-grid coil magnitudes','coil_magnitude',507,428,51,51)
text('Map ellipsis','…',569,435,16,24,16)
text('Phase map title','Phase',590,383,55,19,11,color='682285'); eq('Phase map symbol','∠Ŝ꜀',588,400,59,21,18)
image_stack('Actual fixed coil phases','coil_phase',583,428,45,45); picture('Coil phase colorbar','coil_colorbar',(642,420,647,479)); text('Coil colorbar positive','π',648,413,12,15,8); text('Coil colorbar zero','0',648,441,12,15,8); text('Coil colorbar negative','−π',647,467,15,15,8); # Coil stack uses its actual colorbar at the right.
text('Fixed coil fractions note','Fixed complex coil fractions: coil + object phase',341,503,312,22,12.5,color='792498')
text('Coil interpolation note','Interpolate and normalize for finer grids',374,523,268,18,10.5,color='792498')
path('Calibration into physics',[(320,508),(320,566)],PURPLE,1.7,True)
# Sampling chain.
text('Gaussian initialization label','Gaussian\ninitialization',680,344,101,39,14)
picture('Gaussian initialization — seeded array','initial_noise',(689,392,771,476))
path('Initial noise to frozen U-Net',[(775,437),(794,437)],BLUE,1.5,True)
text('Frozen denoiser label','Frozen denoiser',791,349,148,25,14,True,'155084')
eq('Frozen denoiser formula','Dθ',833,370,57,25,23)
unet('Frozen denoiser U-Net',796,396,130,80)
path('Frozen U-Net to clean estimate',[(931,437),(952,437)],BLUE,1.4,True)
text('Clean estimate title','Clean estimate',949,345,110,22,14); eq('Clean estimate symbol','x̂₀',980,365,49,24,22)
picture('Representative clean estimate — checkpoint evaluation','sampling_clean_estimate',(955,390,1048,479))
path('Clean estimate to consistency',[(1049,437),(1067,437)],BLUE,1.4,True)
box('Data consistency frame',1069,406,95,62,'FFF0D8',ORANGE,1.6,True)
text('Data consistency label','Data\nconsistency',1074,415,86,44,15,False,'A94309')
path('Consistency to update',[(1166,437),(1185,437)],BLUE,1.4,True)
box('Diffusion update frame',1187,406,100,62,'E2F2FD',BLUE,1.5,True)
text('Diffusion update label','Diffusion\nupdate',1193,415,88,44,15)
path('Update to reconstruction',[(1290,437),(1311,437)],BLUE,1.4,True)
text('Reconstruction label','Reconstruction',1307,346,115,24,14); eq('Reconstruction symbol','x̂',1341,367,46,23,22)
picture('Reconstruction — saved baseline result','reconstruction',(1315,390,1415,482))
# Native freeform smooth iteration arrow (editable cubic Bézier segments).
loop=path('Sampling iteration loop',[(1232,472),(1232,484),(1208,513),(909,513),(881,480)],BLUE,1.7,True)
p=loop._element.xpath('.//a:path')[0]
for child in list(p)[1:]: p.remove(child)
# Local coordinates follow the freeform's x/y origin at (881,472).
for coords in [[(351,33),(343,41),(317,41)],[(247,41),(109,41),(31,41)],[(10,41),(0,25),(0,8)]]:
    el=OxmlElement('a:cubicBezTo')
    for x,y in coords:
        pt=OxmlElement('a:pt'); pt.set('x',str(x)); pt.set('y',str(y)); el.append(pt)
    p.append(el)
text('Sampling iteration note','Repeat as σ decreases',970,516,215,23,13,color='175989',italic=True)
path('Physics into data consistency',[(967,573),(1117,573),(1117,471)],ORANGE,1.8,True)
# Large physics explanation inset.
box('Physics inset frame',21,568,945,275,'FFFFFF',TEAL,2.5,True)
text('Physics inset title','What the SPEN degradation operator does',37,575,816,29,22,True,'0B7279',align='left')
box('Magnitude scaling frame',36,615,116,56,'F7FDFD',TEAL,1.5,True)
eq('Magnitude scaling left','m =',46,631,35,27,20); fraction('Magnitude scaling fraction','(x + 1)','2',86,627,49,15)
text('Magnitude scaling note','Denoiser predicts\nnormalized real\nmagnitude x.\nNonnegative\nmagnitude m = (x + 1)/2.',40,681,141,76,11.8,align='left')
path('Magnitude to coil fractions',[(153,646),(177,646)],TEAL,1.7,True)
box('Coil multiplication frame',179,614,106,57,'F3E8F8',PURPLE,1.5,True); eq('Coil multiplication formula','× Ŝ꜀',201,628,61,30,23)
text('Coil multiplication note','Apply fixed\ncomplex coil\nfractions\n(from measurements)',186,681,122,63,11.8,align='left')
path('Coil fractions to SPEN',[(287,645),(313,645)],TEAL,1.7,True)
box('SPEN encoding frame',314,614,131,51,'E4F6F6',TEAL,1.6,True)
eq('SPEN operator symbol','A_{PE}:',348,619,59,24,20); text('SPEN operator label','SPEN encoding',325,640,110,20,13)
picture('Actual encoding matrix magnitude','encoding_magnitude',(341,680,419,729)); eq('Encoding thumbnail symbol','|A_{PE}|',346,725,69,18,12); text('Scanner trajectory note','Same scanner\ntrajectory',330,742,100,32,12)
path('SPEN to RO',[(447,645),(469,645)],TEAL,1.7,True)
box('RO bandlimit frame',470,614,151,51,'E4F6F6',TEAL,1.6,True)
eq('RO operator symbol','D_{RO}:',515,619,59,24,20); text('RO operator label','RO band limitation',477,640,139,20,13)
box('Retained RO frequency band',530,677,34,42,'B1CDEE')
path('RO frequency vertical axis',[(499,675),(499,722)],'555555',.8); path('RO frequency axis',[(499,722),(599,722)],'555555',.8)
vs=[1,0,2,1,0,3,1,5,4,10,18,23,29,34,39,37,34,29,23,14,5,3,5,2,4,1,3,2,1,1,2,1,0]
path('RO spectrum visual approximation',[(500+i*3,717-v) for i,v in enumerate(vs)],'2057CF',1)
text('RO frequency axis label','RO frequency',506,721,91,17,10); text('Acquired bandwidth note','Retain acquired\nbandwidth',491,741,111,33,12)
path('RO to acquired rows',[(623,645),(646,645)],TEAL,1.7,True)
box('PE masking frame',648,614,139,52,'E4F6F6',TEAL,1.6,True)
eq('PE masking symbol','M:',690,619,52,24,21); text('PE masking label','acquired PE rows',658,640,122,20,13)
box('PE row illustration background',654,676,59,56,'161A1C')
for i in range(14): path('PE acquired or missing row '+str(i),[(655,679+i*3.8),(711,679+i*3.8)],'FFFFFF' if i in [2,3,6,10,11] else '777777',1.2,dash=i>=9)
box('Retained row legend swatch',724,689,21,7,'FFFFFF','333333',.8)
path('Missing row legend swatch',[(724,708),(746,708)],'111111',1.7,dash=True)
text('Retained rows legend','Retained rows',751,685,60,16,8.5,align='left'); text('Missing rows legend','Missing rows',751,701,60,16,8.5,align='left')
text('PE mask note','Full or undersampled PE',643,744,162,22,12)
path('Rows to prediction',[(789,645),(818,645)],TEAL,1.7,True)
box('Predicted data frame',820,612,133,130,'F1FAFC',TEAL,1.6,True)
text('Predicted data label','Predicted complex\ndata',828,617,117,29,11.5)
image_stack('Predicted complex data — four receiver coils','predicted_coil',832,657,102,72); eq('Predicted data formula','F꜀(x)',850,744,76,28,20)
box('Forward operator formula band',37,782,915,50,'EFF8FB','BDE6E9',1.5,True)
eq('Forward operator formula left','F꜀(x) = M D_{RO} A_{PE} [Ŝ꜀',172,788,209,36,22)
fraction('Forward operator scaling','x + 1','2',382,785,50,16)
eq('Forward operator closing bracket',']',433,788,24,37,31)
path('Formula example separator',[(616,793),(616,821)],'759DAC',1.3)
text('Acquisition grid example','Example: 192 × 192 image grid → 96 × 96 acquisition',636,793,306,28,11.5,color='174F73')
# Optional phase refinement inset.
box('Phase refinement inset',1045,592,376,254,'FFFFFF',PURPLE,1.6,True,True)
text('Phase refinement title','Experimental extension: phase refinement',1062,600,348,28,15.5,True,'792498')
path('Optional phase refinement link',[(968,597),(1008,645),(1043,645)],PURPLE,1.5,True,True)
text('Optional phase label','Optional',991,607,56,24,11,color='792498')
picture('Actual acquired-even residual phase','residual_phase',(1065,650,1167,701)); picture('Residual phase generated colorbar','residual_colorbar',(1179,642,1191,716)); box('Residual phase map boundary',1065,650,102,51,None,'909090',.6); text('Residual phase domain','Even PE × RO',1064,705,104,17,10)
text('Residual phase upper scale',f"+{RENDER_META['residual_phase_limit_rad']:.1f}",1193,635,29,20,10); text('Residual phase zero scale','0',1193,669,23,20,10); text('Residual phase lower scale',f"−{RENDER_META['residual_phase_limit_rad']:.1f}",1193,705,29,20,10)
text('Residual phase caption','Residual φ [rad]\n48 × 96',1062,727,111,31,11)
text('Phase fitting note','• Fit residual odd/even phase\n   from measured data',1224,637,184,34,12,align='left')
text('Phase row factors','• Even rows: exp(iφ); odd rows: 1',1224,674,186,25,11.3,align='left')
text('Phase optimization note','• Update phase during sampling;\n   prior stays frozen',1224,697,186,34,11.5,align='left')
box('Phase refinement equation frame',1063,770,341,44,'FCF7FD','C69ADA',1.2,True)
eq('Phase refinement equation','F_{φ,c}(x) = M{ exp(iEφ) ⊙ D_{RO} A_{PE} [Ŝ꜀ m] }',1070,774,327,35,20)
text('Even row insertion note','E inserts phase on original even scanner rows BEFORE PE masking.',1061,816,348,23,10.8,color='792498')
text('Legend key','Key:',21,871,34,22,11,True,align='left')
text('Legend prior','Blue:  learned prior',64,871,125,22,12,color='12529C',align='left')
text('Legend physics','Teal:  fixed physics',199,871,125,22,12,color='087F87',align='left')
text('Legend maps','Purple:  phase / coil estimates',330,871,183,22,12,color='792498',align='left')
text('Legend consistency','Orange:  measurement consistency',521,871,254,22,12,color='CD660F',align='left')
# Explicit local styling avoids theme shadows.
for style in list(S._element.xpath('.//p:style')):
    style.getparent().remove(style)
R.save(OUT/'fig1_v2.pptx')
print(OUT/'fig1_v2.pptx'); print('Page-level objects:',len(S.shapes))
